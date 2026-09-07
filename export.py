"""
Budowanie plikow wyjsciowych (xlsx) z gotowych kandydatow do linkowania:
  1. Plik "do oceny" - kandydaci + arkusze pomocnicze do recznej weryfikacji:
     Kandydaci_linkowania, L1_do_uzupelnienia, L2_pod_L1_bez_siostr,
     Pominiete_zbyt_glebokie (kandydaci odcieci limitem max_level_diff),
     Marka_wykluczona_generyczna (kategorie z generycznym leafem, np.
     "Akcesoria", pominiete przy dopasowaniu marka 1-segmentowe), Diagnostyka.
  2. Macierz "do Contentful" - jeden wiersz na zrodlowy URL, w kolejnych
     kolumnach URL-e, do ktorych ten URL ma linkowac.
"""

from __future__ import annotations

import io
from collections import Counter, defaultdict

import openpyxl
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from linking_engine import rule_sort_key


def _autosize(ws):
    for col_cells in ws.columns:
        length = max((len(str(c.value)) if c.value is not None else 0) for c in col_cells)
        ws.column_dimensions[get_column_letter(col_cells[0].column)].width = min(max(length + 2, 10), 80)


def _write_table(ws, headers, rows):
    ws.append(headers)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="4472C4")
    for row in rows:
        ws.append([row.get(h, "") for h in headers])
    ws.freeze_panes = "A2"
    _autosize(ws)


def build_review_workbook(
    all_candidates: list[dict],
    l1_categories: list,
    l2_under_l1_no_siblings: list,
    no_base_found: list,
    pages: list,
    distinct_filtry_main_categories: int | None = None,
    cut_by_depth_candidates: list[dict] | None = None,
    max_level_diff: int | None = None,
    brand_generic_excluded: list | None = None,
) -> bytes:
    cut_by_depth_candidates = cut_by_depth_candidates or []
    brand_generic_excluded = brand_generic_excluded or []

    type_counts = Counter(p.url_type for p in pages)
    rule_counts = Counter()
    for c in all_candidates:
        for r in c["Rule"].split(" + "):
            rule_counts[r] += 1

    wb = openpyxl.Workbook()

    ws1 = wb.active
    ws1.title = "Kandydaci_linkowania"
    headers = [
        "Source_URL", "Target_URL", "Rule", "Source_Level",
        "Source_Type", "Target_Type", "Target_Level",
        "Poziom_roznica", "Anchor",
    ]
    _write_table(ws1, headers, all_candidates)

    ws2 = wb.create_sheet("L1_do_uzupelnienia")
    l1_headers = ["Source_URL", "Source_Type", "Source_Level", "Target_URL", "Target_Type", "Anchor", "Uwaga"]
    l1_rows = [
        {
            "Source_URL": p.url,
            "Source_Type": "category",
            "Source_Level": p.level_label,
            "Target_URL": "",
            "Target_Type": "",
            "Anchor": "",
            "Uwaga": "BRAK REGULY - departament najwyzszego poziomu (brak rodzica w breadcrumbie), "
                     "brak automatycznego 'sasiada tego samego poziomu' - wymaga recznego wskazania "
                     "kategorii komplementarnych/podobnych (linkowanie 'w dol' od tej kategorii jest "
                     "juz w zakladce Kandydaci_linkowania jak dla kazdej innej kategorii)",
        }
        for p in sorted(l1_categories, key=lambda x: x.url)
    ]
    _write_table(ws2, l1_headers, l1_rows)

    ws2b = wb.create_sheet("L2_pod_L1_bez_siostr")
    l2_headers = ["Source_URL", "Source_Type", "Source_Level", "Dzial_L1_rodzic", "Target_URL", "Target_Type", "Anchor", "Uwaga"]
    l2_rows = [
        {
            "Source_URL": p.url,
            "Source_Type": "category",
            "Source_Level": p.level_label,
            "Dzial_L1_rodzic": p.direct_parent_url,
            "Target_URL": "",
            "Target_Type": "",
            "Anchor": "",
            "Uwaga": "BRAK REGULY - rodzic to szeroki dzial L1, wiec inne kategorie pod tym samym "
                     "L1 NIE sa automatycznie traktowane jako spokrewnione 'siostry' - wymaga "
                     "recznej selekcji, jesli w ogole warto tu linkowac",
        }
        for p in sorted(l2_under_l1_no_siblings, key=lambda x: (x.direct_parent_url or "", x.url))
    ]
    _write_table(ws2b, l2_headers, l2_rows)

    ws2c = wb.create_sheet("Pominiete_zbyt_glebokie")
    depth_headers = [
        "Source_URL", "Target_URL", "Rule", "Source_Level",
        "Source_Type", "Target_Type", "Target_Level",
        "Poziom_roznica", "Anchor",
    ]
    _write_table(ws2c, depth_headers, cut_by_depth_candidates)
    if ws2c.max_row == 1:
        ws2c.append(["(brak - wszyscy kandydaci miesca sie w limicie glebokosci)"])

    ws2d = wb.create_sheet("Marka_wykluczona_generyczna")
    generic_headers = ["Source_URL", "Source_Type", "Source_Level", "Anchor", "Uwaga"]
    generic_rows = [
        {
            "Source_URL": p.url,
            "Source_Type": "category",
            "Source_Level": p.level_label,
            "Anchor": p.h1,
            "Uwaga": "BRAK AUTOMATYCZNEJ PROPOZYCJI MARKA (1-segmentowe) - ostatni segment "
                     "breadcrumba jest slowem w pelni generycznym (np. \"Akcesoria\"), samo w "
                     "sobie nie niesie informacji o produkcie. Kategoria dalej dostaje normalnie "
                     "linki z regul kategoria_podrzedna / kategoria_tego_samego_poziomu / filtr_* "
                     "oraz z precyzyjnego dopasowania marka 2-segmentowego, jesli pasuje - traci "
                     "TYLKO orientacyjne dopasowanie marki po samej nazwie. Do recznej oceny, "
                     "jesli warto tu dodac link do marki.",
        }
        for p in sorted(brand_generic_excluded, key=lambda x: x.url)
    ]
    _write_table(ws2d, generic_headers, generic_rows)

    ws3 = wb.create_sheet("Diagnostyka")
    ws3.append(["Metryka", "Wartosc"])
    for c in ws3[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="4472C4")
    diag_rows = [
        ("Wygenerowano kandydatow - unikalne pary source->target (po scaleniu duplikatow)", len(all_candidates)),
        ("", ""),
        ("Rozbicie wedlug reguly (wiersz moze liczyc sie w >1 regule):", ""),
    ]
    for rule, cnt in sorted(rule_counts.items(), key=lambda x: -x[1]):
        diag_rows.append((f"  - {rule}", cnt))
    diag_rows += [
        ("", ""),
        ("Kategorie L1 (brak automatycznych sasiadow tego samego poziomu)", len(l1_categories)),
        ("Kategorie L2 pod L1 (brak automatycznych 'siostr' - szeroki dzial)", len(l2_under_l1_no_siblings)),
        ("Strony filtrowane bez dopasowanej kategorii bazowej po breadcrumbie", len(no_base_found)),
        ("", ""),
        ("Limit roznicy poziomow dla kategoria_podrzedna / filtr_podrzedny (Poziom_roznica)", max_level_diff),
        ("Kandydaci odcieci limitem glebokosci (patrz arkusz Pominiete_zbyt_glebokie)", len(cut_by_depth_candidates)),
        ("Kategorie wykluczone z dopasowania marka 1-segmentowe - nazwa generyczna "
         "(patrz arkusz Marka_wykluczona_generyczna)", len(brand_generic_excluded)),
        ("", ""),
        ("Strony wziete pod uwage razem (po filtrach 3xx/4xx/noindex)", len(pages)),
        ("  - typu category", type_counts.get("category", 0)),
        ("  - typu brand", type_counts.get("brand", 0)),
        ("  - typu filtered_category", type_counts.get("filtered_category", 0)),
    ]
    for label, val in diag_rows:
        ws3.append([label, val])
    if no_base_found:
        ws3.append(["", ""])
        ws3.append(["Strony filtrowane bez dopasowanej kategorii bazowej (do sprawdzenia):", ""])
        for u in no_base_found:
            ws3.append([u, ""])
    _autosize(ws3)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def build_contentful_matrix(candidate_rows: list[dict], max_links: int | None = None) -> bytes:
    """
    candidate_rows: lista dictow z co najmniej kluczami Source_URL, Target_URL
    (opcjonalnie Poziom_roznica, Rule - uzywane do sortowania kolejnosci linkow
    w obrebie jednego zrodla: najpierw wg priorytetu reguly - patrz
    linking_engine.RULE_SORT_ORDER (kategoria_podrzedna, filtr_wlasny,
    kategoria_tego_samego_poziomu, filtr_tego_samego_poziomu, potem reszta),
    potem Poziom_roznica, na koniec Target_URL).

    Zwraca xlsx: pierwsza kolumna Source_URL, kolejne Link_1, Link_2, ...
    """
    by_source: dict[str, list[dict]] = defaultdict(list)
    for r in candidate_rows:
        src = r.get("Source_URL")
        tgt = r.get("Target_URL")
        if not src or not tgt:
            continue
        by_source[src].append(r)

    def sort_key(r):
        return (
            rule_sort_key(str(r.get("Rule") or "")),
            r.get("Poziom_roznica") if isinstance(r.get("Poziom_roznica"), (int, float)) else 0,
            str(r.get("Target_URL") or ""),
        )

    rows_out = []
    max_count = 0
    for src, items in by_source.items():
        items_sorted = sorted(items, key=sort_key)
        targets = []
        seen = set()
        for it in items_sorted:
            t = it["Target_URL"]
            if t in seen:
                continue
            seen.add(t)
            targets.append(t)
        if max_links is not None:
            targets = targets[:max_links]
        max_count = max(max_count, len(targets))
        rows_out.append((src, targets))

    headers = ["Source_URL"] + [f"Link_{i+1}" for i in range(max_count)]

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Matryca_Contentful"
    ws.append(headers)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="4472C4")
    for src, targets in sorted(rows_out, key=lambda x: x[0]):
        ws.append([src] + targets)
    ws.freeze_panes = "A2"
    _autosize(ws)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
